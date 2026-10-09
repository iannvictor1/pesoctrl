import unittest
import os
import uuid
from io import BytesIO
from openpyxl import load_workbook
from concurrent.futures import ThreadPoolExecutor

from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool
from sqlalchemy.schema import CreateSchema, DropSchema

from backend.database import Base, get_db
from backend.inventario import router
from backend.models import InventarioCaixa, EstoqueAgregado


class InventarioTest(unittest.TestCase):
    def setUp(self):
        self.schema = None
        if os.environ.get('INVENTARIO_TEST_POSTGRES') == '1':
            from backend.database import engine
            self.schema = 'teste_inventario_' + uuid.uuid4().hex
            self.base_engine = engine
            with engine.begin() as conn:
                conn.execute(CreateSchema(self.schema))
            self.engine = engine.execution_options(schema_translate_map={None: self.schema})
        else:
            self.engine = create_engine('sqlite://', connect_args={'check_same_thread': False}, poolclass=StaticPool)
        Base.metadata.create_all(self.engine)
        app = FastAPI()
        app.include_router(router)

        def session():
            with Session(self.engine, expire_on_commit=False) as db:
                yield db
        app.dependency_overrides[get_db] = session
        self.client = TestClient(app)

    def tearDown(self):
        self.client.close()
        if self.schema:
            with self.base_engine.begin() as conn:
                conn.execute(DropSchema(self.schema, cascade=True))
        self.engine.dispose()

    def test_operador_obrigatorio_e_persistido(self):
        self.assertEqual(self.client.post('/api/inventarios').status_code, 422)
        for payload in [{}, {'operador_nome': None}, {'operador_nome': ''},
                        {'operador_nome': '   '}, {'operador_nome': 'a' * 121}]:
            self.assertEqual(self.client.post('/api/inventarios', json=payload).status_code, 422)
        self.assertEqual(self.client.get('/api/inventarios').json(), [])
        response = self.client.post('/api/inventarios', json={'operador_nome': '  João Silva  '})
        self.assertEqual(response.status_code, 200)
        inv = response.json()['id']
        for dados in [response.json(), self.client.get(f'/api/inventarios/{inv}').json(),
                      self.client.get('/api/inventarios').json()[0],
                      self.client.get('/api/inventarios/resumos').json()[0],
                      self.client.post(f'/api/inventarios/{inv}/encerrar').json()]:
            self.assertEqual(dados['operador_nome'], 'João Silva')
        wb = load_workbook(BytesIO(self.client.get(f'/api/inventarios/{inv}/exportar').content))
        self.assertEqual(wb['Resumo']['B10'].value, 'João Silva')

    def test_migracao_preserva_inventario_antigo(self):
        from unittest.mock import patch
        from sqlalchemy import text
        from backend.database import ensure_runtime_schema

        engine = create_engine('sqlite://')
        try:
            with engine.begin() as conn:
                conn.execute(text('CREATE TABLE caixas (id INTEGER, codigo_barras_original TEXT)'))
                conn.execute(text('CREATE TABLE inventarios (id INTEGER PRIMARY KEY, status VARCHAR(20))'))
                conn.execute(text("INSERT INTO inventarios (id, status) VALUES (1, 'encerrado')"))
            with patch('backend.database.engine', engine):
                ensure_runtime_schema()
                ensure_runtime_schema()
            with engine.connect() as conn:
                self.assertEqual(tuple(conn.execute(text('SELECT id, status, operador_nome FROM inventarios')).one()), (1, 'encerrado', None))
        finally:
            engine.dispose()

    @unittest.skipUnless(os.environ.get('INVENTARIO_TEST_POSTGRES') == '1', 'Concorrência requer PostgreSQL')
    def test_coletores_simultaneos(self):
        inv, pallet, path = self.criar()
        outro = self.novo_pallet(inv)
        inv2, pallet2, path2 = self.criar()
        destinos = [path, f'/api/inventarios/{inv}/pallets/{outro}', path2]

        def bipar(n):
            return self.client.post(destinos[n % 3] + '/caixas', json=dict(
                codigo='ETIQUETA-REPETIDA', descricao='Produto concorrente',
                validade='2027-01-31', peso_kg='10.125', leitura_id=f'simultanea-{n}',
                coletor_nome=f'Coletor {n % 3 + 1}',
            ))

        with ThreadPoolExecutor(max_workers=6) as pool:
            respostas = list(pool.map(bipar, range(30)))
        self.assertTrue(all(r.status_code == 200 for r in respostas), [r.text for r in respostas if r.status_code != 200])
        with ThreadPoolExecutor(max_workers=3) as pool:
            repetidas = list(pool.map(bipar, [0, 0, 0]))
        self.assertEqual(len({r.json()['id'] for r in repetidas}), 1)
        dados = self.client.get(f'/api/inventarios/{inv}').json()
        dados2 = self.client.get(f'/api/inventarios/{inv2}').json()
        for p in dados['pallets'] + dados2['pallets']:
            self.assertEqual(len(p['caixas']), 10)
            self.assertEqual(len({c['coletor_nome'] for c in p['caixas']}), 1)
            self.assertEqual(sum(float(c['peso_kg']) for c in p['caixas']), 101.25)

    def novo_pallet(self, inv):
        resposta = self.client.post(f'/api/inventarios/{inv}/pallets', json={'unitizador': '000123'})
        self.assertEqual(resposta.status_code, 200, resposta.text)
        pallet = resposta.json()['id']
        resposta = self.client.post(f'/api/inventarios/{inv}/pallets/{pallet}/endereco', json={'codigo': '1959947609'})
        self.assertEqual(resposta.status_code, 200, resposta.text)
        return pallet

    def test_fluxo_unitizador_endereco_e_caixas(self):
        inv = self.client.post('/api/inventarios', json={'operador_nome': 'Maria Silva'}).json()['id']
        base = f'/api/inventarios/{inv}/pallets'
        self.assertEqual(self.client.post(base).status_code, 422)
        self.assertEqual(self.client.post(base, json={'unitizador': '   '}).status_code, 422)
        primeiro = self.client.post(base, json={'unitizador': '0000123'}).json()
        segundo = self.client.post(base, json={'unitizador': '0000123'}).json()
        self.assertNotEqual(primeiro['id'], segundo['id'])
        self.assertEqual(primeiro['unitizador'], '0000123')
        path = f"{base}/{primeiro['id']}"
        caixa = dict(codigo='CX', descricao='Produto', validade='2027-01-31', peso_kg='1.250')
        self.assertEqual(self.client.post(path + '/caixas', json=caixa).status_code, 409)
        self.assertEqual(self.client.post(path + '/endereco', json={'codigo': 'inexistente'}).status_code, 422)
        self.assertEqual(self.client.post(path + '/caixas', json=caixa).status_code, 409)
        endereco = self.client.post(path + '/endereco', json={'codigo': ' 2015154011\r\n'}).json()
        self.assertEqual(endereco['endereco'], 'R 1 ; C 1 ; A 3 ; AP 2')
        self.assertEqual(endereco['endereco_codigo'], '2015154011')
        self.assertEqual(self.client.post(path + '/caixas', json=caixa).status_code, 200)
        self.assertEqual(self.client.post(path + '/endereco', json={'codigo': '2015154011'}).status_code, 200)
        self.assertEqual(self.client.post(path + '/endereco', json={'codigo': '1959947609'}).status_code, 409)
        dados = self.client.get(f'/api/inventarios/{inv}').json()['pallets']
        self.assertEqual(dados[0]['unitizador'], '0000123')
        self.assertEqual(dados[0]['endereco'], endereco['endereco'])
        self.assertIsNone(dados[1]['endereco'])
        self.assertEqual(dados[1]['caixas'], [])
        self.client.post(path + '/encerrar')
        self.assertEqual(self.client.post(path + '/endereco', json={'codigo': '2015154011'}).status_code, 409)

    def test_migracao_preserva_pallet_antigo(self):
        from unittest.mock import patch
        from sqlalchemy import text, inspect
        from backend.database import ensure_runtime_schema
        engine = create_engine('sqlite://')
        try:
            with engine.begin() as conn:
                conn.execute(text('CREATE TABLE caixas (id INTEGER PRIMARY KEY)'))
                conn.execute(text('CREATE TABLE inventario_pallets (id INTEGER PRIMARY KEY)'))
                conn.execute(text('INSERT INTO inventario_pallets (id) VALUES (1)'))
            with patch('backend.database.engine', engine):
                ensure_runtime_schema()
                ensure_runtime_schema()
            self.assertTrue({'unitizador', 'endereco_codigo', 'endereco'}.issubset(
                {c['name'] for c in inspect(engine).get_columns('inventario_pallets')}))
            with engine.connect() as conn:
                self.assertEqual(conn.execute(text('SELECT id, unitizador, endereco FROM inventario_pallets')).one(), (1, None, None))
        finally:
            engine.dispose()

    def criar(self):
        inv = self.client.post('/api/inventarios', json={'operador_nome': 'Maria Silva'}).json()['id']
        pallet = self.novo_pallet(inv)
        return inv, pallet, f'/api/inventarios/{inv}/pallets/{pallet}'

    def test_separacao_persistencia_duplicidade_e_encerramento(self):
        inv, pallet, path = self.criar()
        inv2, pallet2, path2 = self.criar()
        payload = dict(codigo='CX-TESTE', descricao='Produto teste', validade='2027-01-31', peso_kg='12.345', leitura_id='leitura-1')
        self.assertEqual(self.client.post(path+'/caixas', json=payload).status_code, 200)
        self.assertEqual(self.client.post(path+'/caixas', json=payload).status_code, 200)
        for n in range(2, 6):
            self.assertEqual(self.client.post(path+'/caixas', json={**payload, 'leitura_id': f'leitura-{n}'}).status_code, 200)
        self.assertEqual(self.client.post(f'/api/inventarios/{inv}/pallets/{pallet2}/caixas', json=payload).status_code, 404)
        self.assertEqual(self.client.post(path2+'/caixas', json=payload).status_code, 200)
        with Session(self.engine) as db:
            caixas = db.scalars(select(InventarioCaixa)).all()
            self.assertEqual(len(caixas), 6)
            self.assertEqual(str(caixas[0].peso_kg), '12.345')
            self.assertEqual(caixas[0].validade, '31/01/2027')
        self.assertEqual(self.client.post(f'/api/inventarios/{inv}/encerrar').status_code, 409)
        self.assertEqual(self.client.post(path+'/encerrar').status_code, 200)
        self.assertEqual(self.client.post(path+'/caixas', json={**payload, 'codigo': 'OUTRO'}).status_code, 409)
        self.assertEqual(self.client.post(f'/api/inventarios/{inv}/encerrar').status_code, 200)
        self.assertEqual(self.client.post(f'/api/inventarios/{inv}/pallets', json={'unitizador': '001'}).status_code, 409)
        self.assertEqual(len(self.client.get(f'/api/inventarios/{inv2}').json()['pallets'][0]['caixas']), 1)

    def test_etiqueta_automatica_e_dados_invalidos(self):
        inv, pallet, path = self.criar()
        with Session(self.engine) as db:
            db.add(EstoqueAgregado(importacao_id='teste', produto_codigo='123', produto_descricao='Carne teste', saldo_disponivel_atual_kg=100))
            db.commit()
        resposta = self.client.post(path+'/caixas', json={'codigo': 'LA,001233712.34515270131'})
        self.assertEqual(resposta.status_code, 200, resposta.text)
        self.assertEqual(resposta.json()['descricao'], 'Carne teste')
        self.assertEqual(resposta.json()['validade'], '31/01/2027')
        self.assertEqual(self.client.post(path+'/caixas', json={'codigo': 'DESCONHECIDO'}).status_code, 422)
        payload = dict(codigo='OUTRO', descricao='Produto', validade='31/02/2027', peso_kg='10')
        self.assertEqual(self.client.post(path+'/caixas', json=payload).status_code, 422)
        self.assertEqual(self.client.post(path+'/caixas', json={**payload, 'peso_kg': '-1'}).status_code, 422)
        with Session(self.engine) as db:
            self.assertEqual(db.scalar(select(EstoqueAgregado.saldo_disponivel_atual_kg)), 100)

    def test_resumos_pallets_sem_multiplicar_caixas(self):
        inv, pallet, path = self.criar()
        vazio = self.client.post('/api/inventarios', json={'operador_nome': 'Maria Silva'}).json()['id']
        outro = self.novo_pallet(inv)
        for numero, destino in enumerate([path, path, f'/api/inventarios/{inv}/pallets/{outro}']):
            response = self.client.post(destino + '/caixas', json=dict(
                codigo='REPETIDO', descricao='Produto', validade='2027-01-31',
                peso_kg='12.345', leitura_id=f'resumo-{numero}',
            ))
            self.assertEqual(response.status_code, 200)
        response = self.client.get('/api/inventarios/resumos')
        self.assertEqual(response.status_code, 200)
        resumos = {r['id']: r for r in response.json()}
        self.assertEqual(resumos[inv]['total_pallets'], 2)
        self.assertEqual(resumos[inv]['total_caixas'], 3)
        self.assertEqual(float(resumos[inv]['peso_total_kg']), 37.035)
        self.assertEqual(resumos[vazio]['total_pallets'], 0)
        self.assertEqual(resumos[vazio]['total_caixas'], 0)
        self.assertEqual(float(resumos[vazio]['peso_total_kg']), 0)

    def test_exportacao_excel(self):
        inv, pallet, path = self.criar()
        self.criar()  # Outro inventário não deve aparecer na exportação.
        for n in range(2):
            self.client.post(path + '/caixas', json=dict(
                codigo='00012345678901234567890', descricao='=PRODUTO',
                validade='2027-01-31', peso_kg='12.345', coletor_nome='=COLETOR', leitura_id=f'excel-{n}',
            ))
        self.client.post(path + '/encerrar')
        self.client.post(f'/api/inventarios/{inv}/encerrar')
        response = self.client.get(f'/api/inventarios/{inv}/exportar')
        self.assertEqual(response.status_code, 200)
        self.assertIn(f'inventario_{inv}.xlsx', response.headers['content-disposition'])
        wb = load_workbook(BytesIO(response.content))
        self.assertEqual(wb.sheetnames, ['Resumo', 'Pallets', 'Caixas'])
        self.assertEqual(wb['Resumo']['B2'].value, 'encerrado')
        self.assertEqual(wb['Resumo']['B7'].value, 2)
        self.assertEqual(wb['Resumo']['B8'].value, 24.69)
        self.assertEqual(wb['Pallets']['G2'].value, 24.69)
        self.assertEqual(wb['Caixas'].max_row, 3)
        self.assertEqual(wb['Caixas']['A2'].value, inv)
        self.assertEqual(wb['Caixas']['B2'].value, pallet)
        self.assertEqual(wb['Caixas']['D2'].value, '=PRODUTO')
        self.assertEqual(wb['Caixas']['D2'].data_type, 's')
        self.assertEqual(wb['Caixas']['G2'].data_type, 's')
        self.assertEqual(wb['Caixas']['F2'].value, 12.345)
        self.assertEqual(wb['Pallets']['I2'].value, '000123')
        self.assertEqual(wb['Caixas']['I2'].value, '000123')
        self.assertEqual(wb['Caixas']['J2'].value, 'R 1 ; C 1 ; A 1 ; AP 1')
        self.assertEqual(wb['Caixas']['K2'].value, '1959947609')
        vazio = self.client.post('/api/inventarios', json={'operador_nome': 'Maria Silva'}).json()['id']
        response = self.client.get(f'/api/inventarios/{vazio}/exportar')
        wb = load_workbook(BytesIO(response.content))
        self.assertEqual(wb['Resumo']['B8'].value, 0)
        self.assertEqual(wb['Caixas'].max_row, 1)
        self.assertEqual(self.client.get('/api/inventarios/999999/exportar').status_code, 404)


if __name__ == '__main__':
    unittest.main()
