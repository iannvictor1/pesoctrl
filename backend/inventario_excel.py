from datetime import datetime
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


def gerar_planilha(inventario: dict) -> bytes:
    wb = Workbook()
    resumo = wb.active
    resumo.title = "Resumo"
    pallets = inventario["pallets"]
    caixas = [c for p in pallets for c in p["caixas"]]
    peso_total = sum((c["peso_kg"] for c in caixas), Decimal("0"))

    def linha(ws, valores):
        ws.append(valores)
        for cell in ws[ws.max_row]:
            # Etiquetas e nomes são texto, mesmo quando começam com '='.
            if isinstance(cell.value, str):
                cell.data_type = "s"
            if isinstance(cell.value, datetime):
                cell.number_format = "dd/mm/yyyy hh:mm:ss"

    linha(resumo, [f"Inventário {inventario['id']}", "Informações"])
    for row in [
        ["Status", inventario["status"]],
        ["Início", inventario["inicio"]],
        ["Encerramento", inventario["fim"]],
        ["Exportado em", datetime.now()],
        ["Total de pallets", len(pallets)],
        ["Total de caixas", len(caixas)],
        ["Peso total (kg)", peso_total],
        ["Observação", "Posição registrada no momento da exportação."],
        ["Operador", inventario.get("operador_nome") or "Não informado"],
    ]:
        linha(resumo, row)
    resumo["B8"].number_format = '#,##0.000'

    ws_pallets = wb.create_sheet("Pallets")
    linha(ws_pallets, ["Inventário", "Pallet", "Status", "Início", "Encerramento", "Caixas", "Peso total (kg)", "Coletores", "Unitizador", "Endereço", "Código do endereço"])
    ws_caixas = wb.create_sheet("Caixas")
    linha(ws_caixas, ["Inventário", "Pallet", "ID da caixa", "Produto", "Validade", "Peso (kg)", "Coletor", "Data da leitura", "Unitizador", "Endereço", "Código do endereço"])
    for p in pallets:
        linha(ws_pallets, [inventario["id"], p["id"], p["status"], p["inicio"], p["fim"], len(p["caixas"]),
                          sum((c["peso_kg"] for c in p["caixas"]), Decimal("0")),
                          ", ".join(sorted({c["coletor_nome"] for c in p["caixas"] if c.get("coletor_nome")})), p.get("unitizador"), p.get("endereco"), p.get("endereco_codigo")])
        ws_pallets.cell(ws_pallets.max_row, 7).number_format = '#,##0.000'
        for c in sorted(p["caixas"], key=lambda c: c["id"]):
            linha(ws_caixas, [inventario["id"], p["id"], c["id"], c["descricao"], c["validade"], c["peso_kg"], c.get("coletor_nome"), c["data_hora"], p.get("unitizador"), p.get("endereco"), p.get("endereco_codigo")])
            ws_caixas.cell(ws_caixas.max_row, 6).number_format = '#,##0.000'

    for ws, widths in [(resumo, [28, 64]), (ws_pallets, [14, 12, 16, 23, 23, 12, 20, 38, 24, 40, 24]), (ws_caixas, [14, 12, 15, 50, 16, 18, 28, 23, 24, 40, 24])]:
        ws.freeze_panes = "A2"
        ws.sheet_view.showGridLines = False
        ws.row_dimensions[1].height = 28
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor="183B36")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(vertical="center")
        for index, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(index)].width = width
        if ws != resumo:
            ws.auto_filter.ref = ws.dimensions
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if cell.row % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor="EDF4F1")
    stream = BytesIO()
    wb.save(stream)
    return stream.getvalue()
