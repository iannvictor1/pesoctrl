from __future__ import annotations

from datetime import datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
import base64
import re
import uuid

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from backend.models import Caixa, EstoqueAgregado


def decimal_kg(value) -> Decimal:
    if value is None or value == "":
        return Decimal("0.000")
    try:
        return Decimal(str(value).replace(",", ".")).quantize(Decimal("0.001"))
    except (InvalidOperation, ValueError):
        return Decimal("0.000")


def formatar_validade(value) -> str | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.strftime("%d/%m/%Y")
    texto = str(value).strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(texto[:10], fmt).strftime("%d/%m/%Y")
        except ValueError:
            pass
    return texto


def separar_produto(value) -> tuple[str, str]:
    texto = str(value or "").strip()
    match = re.match(r"^\s*(?P<codigo>\d+)\s*-\s*(?P<descricao>.+)$", texto)
    if not match:
        return "", texto
    return match.group("codigo").lstrip("0") or "0", match.group("descricao").strip()


def importar_estoque_xlsx(db: Session, arquivo_base64: str, nome_arquivo: str | None = None) -> dict:
    conteudo = base64.b64decode(arquivo_base64)
    workbook = load_workbook(BytesIO(conteudo), data_only=True, read_only=True)
    worksheet = workbook[workbook.sheetnames[0]]

    linhas = list(worksheet.iter_rows(values_only=True))
    if not linhas:
        return {"ok": False, "erro": "planilha vazia"}

    importacao_id = datetime.now().strftime("%Y%m%d_%H%M%S") + "_" + uuid.uuid4().hex[:6].upper()
    agora = datetime.now()
    total = 0
    ignoradas = 0
    novos_registros = []

    for row in linhas[1:]:
        if not row or not row[2]:
            ignoradas += 1
            continue

        codigo, descricao = separar_produto(row[2])
        saldo_disponivel = decimal_kg(row[8] if len(row) > 8 else None)
        if not codigo or saldo_disponivel <= 0:
            ignoradas += 1
            continue

        novos_registros.append(
            {
                "importacao_id": importacao_id,
                "arquivo_origem": nome_arquivo,
                "importado_em": agora,
                "deposito": str(row[0]).strip() if len(row) > 0 and row[0] is not None else None,
                "endereco": str(row[1]).strip() if len(row) > 1 and row[1] is not None else None,
                "produto_codigo": codigo,
                "produto_descricao": descricao,
                "unitizador": str(row[3]).strip() if len(row) > 3 and row[3] is not None else None,
                "condicao": str(row[4]).strip() if len(row) > 4 and row[4] is not None else None,
                "situacao": str(row[5]).strip() if len(row) > 5 and row[5] is not None else None,
                "saldo_fisico_kg": decimal_kg(row[6] if len(row) > 6 else None),
                "saldo_reservado_kg": decimal_kg(row[7] if len(row) > 7 else None),
                "saldo_disponivel_inicial_kg": saldo_disponivel,
                "saldo_disponivel_atual_kg": saldo_disponivel,
                "bloqueado_movimentacao_kg": decimal_kg(row[9] if len(row) > 9 else None),
                "validade": formatar_validade(row[10] if len(row) > 10 else None),
                "lote": str(row[11]).strip() if len(row) > 11 and row[11] is not None else None,
            }
        )
        total += 1

    if not novos_registros:
        return {"ok": False, "erro": "nenhuma linha valida encontrada na planilha", "linhas_ignoradas": ignoradas}

    agregados_substituidos = (
        db.query(EstoqueAgregado)
        .filter(EstoqueAgregado.status == "ativo")
        .update({"status": "substituido"}, synchronize_session=False)
    )
    caixas_substituidas = (
        db.query(Caixa)
        .filter(Caixa.status == "em_estoque")
        .update({"status": "substituido_wms"}, synchronize_session=False)
    )
    db.add_all(EstoqueAgregado(**registro) for registro in novos_registros)
    db.commit()
    return {
        "ok": True,
        "importacao_id": importacao_id,
        "linhas_importadas": total,
        "linhas_ignoradas": ignoradas,
        "estoques_substituidos": int(agregados_substituidos or 0),
        "caixas_substituidas": int(caixas_substituidas or 0),
    }


def ultima_importacao_id(db: Session) -> str | None:
    return db.scalar(
        select(EstoqueAgregado.importacao_id)
        .where(EstoqueAgregado.status == "ativo")
        .order_by(EstoqueAgregado.importado_em.desc(), EstoqueAgregado.id.desc())
        .limit(1)
    )
